"""Document formatters generating Markdown, JSON, PDF, and DOCX meeting files."""

import json
from pathlib import Path
from typing import Any

from docx import Document
from reportlab.lib import colors  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import letter  # type: ignore[import-untyped]
from reportlab.lib.styles import (  # type: ignore[import-untyped]
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.platypus import (  # type: ignore[import-untyped]
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


class MarkdownFormatter:
    """Renders meeting data into clean Markdown format."""

    @staticmethod
    def format(data: dict[str, Any]) -> str:
        """Format meeting data payload into Markdown string."""
        title = data.get("title", "Untitled Meeting")
        created_at = data.get("created_at", "")
        duration = data.get("duration", 0.0)

        lines = [
            f"# {title}",
            f"**Date:** {created_at} | **Duration:** {duration:.1f}s\n",
            "---",
            "\n## Executive Summary",
            data.get("summary", {}).get("executive_summary", "No summary generated."),
        ]

        takeaways = data.get("summary", {}).get("key_takeaways", [])
        if takeaways:
            lines.append("\n### Key Takeaways")
            for kt in takeaways:
                lines.append(f"- {kt}")

        action_items = data.get("action_items", [])
        if action_items:
            lines.append("\n## Action Items")
            for ai in action_items:
                assignee = f" (@{ai['assignee']})" if ai.get("assignee") else ""
                due = f" [Due: {ai['due_date']}]" if ai.get("due_date") else ""
                lines.append(f"- [ ] {ai['content']}{assignee}{due}")

        decisions = data.get("decisions", [])
        if decisions:
            lines.append("\n## Decisions")
            for d in decisions:
                lines.append(f"- {d['content']}")

        transcripts = data.get("transcripts", [])
        if transcripts:
            lines.append("\n## Transcript")
            for t in transcripts:
                spk = (
                    f"**{t.get('speaker_name', 'Speaker')}**: "
                    if t.get("speaker_name")
                    else ""
                )
                lines.append(f"- [{t['timestamp']:.1f}s] {spk}{t['text']}")

        return "\n".join(lines)


class JSONFormatter:
    """Renders meeting data into formatted JSON string."""

    @staticmethod
    def format(data: dict[str, Any]) -> str:
        """Format meeting data payload into pretty JSON string."""
        return json.dumps(data, indent=2, ensure_ascii=False)


class PDFFormatter:
    """Renders meeting data into a PDF document using ReportLab."""

    @staticmethod
    def export(data: dict[str, Any], output_path: str | Path) -> None:
        """Generate PDF report document.

        Args:
            data: Meeting data dictionary.
            output_path: Destination file path.
        """
        doc = SimpleDocTemplate(str(output_path), pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            "MeetingTitle",
            parent=styles["Heading1"],
            fontSize=20,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=12,
        )
        heading_style = ParagraphStyle(
            "SectionHeading",
            parent=styles["Heading2"],
            fontSize=14,
            textColor=colors.HexColor("#4F46E5"),
            spaceBefore=12,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "BodyTextCustom", parent=styles["BodyText"], fontSize=10, spaceAfter=4
        )

        # Title & Metadata
        story.append(Paragraph(data.get("title", "Untitled Meeting"), title_style))
        story.append(
            Paragraph(
                f"<b>Date:</b> {data.get('created_at', '')} | "
                f"<b>Duration:</b> {data.get('duration', 0.0):.1f}s",
                body_style,
            )
        )
        story.append(Spacer(1, 10))

        # Executive Summary
        story.append(Paragraph("Executive Summary", heading_style))
        summary_text = data.get("summary", {}).get(
            "executive_summary", "No summary generated."
        )
        story.append(Paragraph(summary_text, body_style))
        story.append(Spacer(1, 10))

        # Action Items Table
        action_items = data.get("action_items", [])
        if action_items:
            story.append(Paragraph("Action Items", heading_style))
            table_data = [["Action Item", "Assignee", "Due Date"]]
            for ai in action_items:
                table_data.append(
                    [
                        ai.get("content", ""),
                        ai.get("assignee") or "-",
                        ai.get("due_date") or "-",
                    ]
                )
            t = Table(table_data, colWidths=[280, 100, 100])
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 10))

        # Transcripts
        transcripts = data.get("transcripts", [])
        if transcripts:
            story.append(Paragraph("Transcript", heading_style))
            for tr in transcripts[:50]:  # Cap PDF preview
                spk = (
                    f"<b>{tr.get('speaker_name', 'Speaker')}:</b> "
                    if tr.get("speaker_name")
                    else ""
                )
                line = f"[{tr['timestamp']:.1f}s] {spk}{tr['text']}"
                story.append(Paragraph(line, body_style))

        doc.build(story)


class DOCXFormatter:
    """Renders meeting data into Microsoft Word (.docx) document."""

    @staticmethod
    def export(data: dict[str, Any], output_path: str | Path) -> None:
        """Generate Word DOCX report document.

        Args:
            data: Meeting data dictionary.
            output_path: Destination file path.
        """
        doc = Document()
        doc.add_heading(data.get("title", "Untitled Meeting"), level=0)

        doc.add_paragraph(
            f"Date: {data.get('created_at', '')} | "
            f"Duration: {data.get('duration', 0.0):.1f}s"
        )

        # Executive Summary
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(
            data.get("summary", {}).get("executive_summary", "No summary generated.")
        )

        # Action Items
        action_items = data.get("action_items", [])
        if action_items:
            doc.add_heading("Action Items", level=1)
            table = doc.add_table(rows=1, cols=3)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Action Item"
            hdr_cells[1].text = "Assignee"
            hdr_cells[2].text = "Due Date"

            for ai in action_items:
                row_cells = table.add_row().cells
                row_cells[0].text = ai.get("content", "")
                row_cells[1].text = ai.get("assignee") or "-"
                row_cells[2].text = ai.get("due_date") or "-"

        # Transcript
        transcripts = data.get("transcripts", [])
        if transcripts:
            doc.add_heading("Transcript", level=1)
            for tr in transcripts:
                spk = (
                    f"{tr.get('speaker_name', 'Speaker')}: "
                    if tr.get("speaker_name")
                    else ""
                )
                p = doc.add_paragraph()
                p.add_run(f"[{tr['timestamp']:.1f}s] ").bold = True
                p.add_run(f"{spk}{tr['text']}")

        doc.save(str(output_path))
